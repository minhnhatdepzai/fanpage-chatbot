"""RAG tài liệu: đọc/làm sạch/chia đoạn, pgvector (lọc public/internal), graph + trợ lý nội bộ, endpoint embeddings."""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest

from app.config import get_settings
from app.conversation.graph import GraphContext, build_graph, run_turn
from app.conversation.knowledge import KnowledgeBase, KnowledgeEntry
from app.conversation.memory_store import InMemoryConversationStore
from app.conversation.profile import FanpageProfile
from app.providers.llm import FakeProvider
from app.rag.embedder import HashingEmbedder
from app.rag.parsing import DocumentError, Page, chunk_pages, clean_text, parse_file


# ------------------------------------------------------------------------------------------ parsing
def test_clean_text_keeps_vietnamese_and_real_hyphens():
    raw = "Thi  tốt\u00a0nghiệp\r\nCOVID-\n19\n\n\n\nHọc sinh"
    assert clean_text(raw) == "Thi tốt nghiệp\nCOVID-19\n\nHọc sinh"
    decomposed = "Vie\u0302\u0323t Nam"  # "Việt" dạng tổ hợp -> NFC
    assert clean_text(decomposed) == "Việt Nam"


def test_chunks_respect_size_overlap_and_page_ranges():
    pages = [Page(1, "Điều 1. Phạm vi điều chỉnh. " * 20), Page(2, "Điều 2. Đối tượng áp dụng. " * 20)]
    chunks = chunk_pages(pages, max_chars=300, overlap_chars=60)
    assert all(len(c.text) <= 300 for c in chunks)
    assert chunks[0].page_from == 1 and chunks[-1].page_to == 2
    assert any(
        c.page_from == 1 and c.page_to == 2 for c in chunks
    )  # đoạn vắt qua 2 trang giữ đúng khoảng trang
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_repeated_headers_and_page_numbers_are_removed(tmp_path: Path):
    body = "\n\n".join(
        f"CÔNG TY ABC - TÀI LIỆU NỘI BỘ\nNội dung riêng của trang {i} về quy trình.\nTrang {i}"
        for i in range(1, 5)
    )
    f = tmp_path / "a.txt"
    f.write_text(body, encoding="utf-8")
    doc = parse_file(f, max_bytes=10**6, max_pages=10)
    text = doc.pages[0].text
    assert "Nội dung riêng của trang 3" in text and "Trang 3" not in text.split("\n")


def test_docx_paragraphs_and_tables_are_extracted(tmp_path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("Quy trình nghỉ phép: gửi đơn trước 3 ngày làm việc.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Loại nghỉ", "Số ngày"
    table.cell(1, 0).text, table.cell(1, 1).text = "Nghỉ phép năm", "12"
    buf = io.BytesIO()
    d.save(buf)
    f = tmp_path / "quy-trinh.docx"
    f.write_bytes(buf.getvalue())
    doc = parse_file(f, max_bytes=10**6, max_pages=10)
    assert "gửi đơn trước 3 ngày" in doc.pages[0].text and "Nghỉ phép năm | 12" in doc.pages[0].text


@pytest.mark.parametrize(
    ("name", "content", "msg"),
    [
        ("a.exe", b"MZ", "chưa hỗ trợ"),
        ("a.pdf", b"not a pdf", "không phải PDF"),
        ("a.docx", b"xx", "không phải Word"),
    ],
)
def test_bad_files_are_rejected(tmp_path: Path, name, content, msg):  # type: ignore[no-untyped-def]
    f = tmp_path / name
    f.write_bytes(content)
    with pytest.raises(DocumentError, match=msg):
        parse_file(f, max_bytes=10**6, max_pages=10)


def test_file_size_limit(tmp_path: Path):
    f = tmp_path / "big.txt"
    f.write_bytes(b"a" * 2000)
    with pytest.raises(DocumentError, match="vượt giới hạn"):
        parse_file(f, max_bytes=1000, max_pages=10)


# ------------------------------------------------------------------------------------------ pgvector
@pytest.mark.db
async def test_ingest_search_visibility_and_delete(db, tmp_path: Path):
    from app.rag import store
    from app.rag.ingest import ingest_file
    from app.rag.retriever import DocSearch

    s = get_settings().model_copy(update={"doc_min_similarity": 0.2})
    emb = HashingEmbedder(dim=s.embedding_dim)
    pub = tmp_path / "hoc-phi.txt"
    pub.write_text("Học sinh trường công lập được miễn học phí từ năm học 2025-2026.", encoding="utf-8")
    internal = tmp_path / "luong.txt"
    internal.write_text("Bảng lương nội bộ: nhân viên thử việc hưởng 85% lương chính thức.", encoding="utf-8")
    r1 = await ingest_file(pub, settings=s, embedder=emb, visibility="public", actor="test")
    await ingest_file(internal, settings=s, embedder=emb, visibility="internal", actor="test")
    assert r1["status"] == "ingested" and r1["chunks"] == 1
    again = await ingest_file(pub, settings=s, embedder=emb, visibility="public", actor="test")
    assert again["status"] == "skipped_duplicate"

    ds = DocSearch(s, emb)
    public_hits = await ds.search("nhân viên thử việc hưởng bao nhiêu lương", visibility=store.PUBLIC)
    assert all("lương" not in h.content for h in public_hits)  # bot công khai KHÔNG thấy tài liệu nội bộ
    internal_hits = await ds.search("nhân viên thử việc hưởng bao nhiêu lương", visibility=store.ALL)
    assert internal_hits and "85%" in internal_hits[0].content
    fee = await ds.search("trường công có miễn học phí không", visibility=store.PUBLIC)
    assert fee and fee[0].source == "Tài liệu: hoc-phi.txt"

    strict = DocSearch(get_settings().model_copy(update={"doc_min_similarity": 0.99}), emb)
    assert await strict.search("trường công có miễn học phí không", visibility=store.PUBLIC) == []

    docs = await store.list_documents()
    assert await store.delete_document(docs[0]["id"], "test")
    assert len(await store.list_documents()) == 1


@pytest.mark.db
async def test_public_ingest_refuses_personal_data(db, tmp_path: Path):
    from app.rag.ingest import ingest_file

    f = tmp_path / "ds.txt"
    f.write_text("Danh sách liên hệ: anh Nam 0912345678, email nam@example.com", encoding="utf-8")
    with pytest.raises(DocumentError, match="dữ liệu cá nhân"):
        await ingest_file(
            f, settings=get_settings(), embedder=HashingEmbedder(), visibility="public", actor="t"
        )


# ------------------------------------------------------------------------------------------ graph + trợ lý
class FakeDocs:
    def __init__(self, entries=None, fail: bool = False):  # type: ignore[no-untyped-def]
        self.entries = entries or []
        self.fail = fail
        self.calls: list[frozenset[str]] = []

    async def search(self, query, *, visibility, top_k=None):  # type: ignore[no-untyped-def]
        self.calls.append(visibility)
        if self.fail:
            raise RuntimeError("db down")
        return self.entries


DOC = KnowledgeEntry(
    "doc:1:1",
    "Quy chế thi tốt nghiệp THPT (trang 12)",
    "Tài liệu: quy-che.pdf",
    "Thí sinh được mang vào phòng thi bút viết, thước kẻ, máy tính cầm tay không có chức năng soạn thảo.",
)


async def _turn(docs, text: str, responder):  # type: ignore[no-untyped-def]
    clock = [datetime(2026, 9, 24, 10, 0, tzinfo=UTC)]
    store = InMemoryConversationStore(now=lambda: clock[0])
    provider = FakeProvider(responder)
    ctx = GraphContext(
        settings=get_settings(),
        store=store,
        llm=provider,
        profile=FanpageProfile(name="AI Test"),
        knowledge=KnowledgeBase(),
        docs=docs,
        now=lambda: clock[0],
    )
    cid = store.ensure_conversation("PAGE1", "u1")
    store.add_user_message(cid, text, ts=clock[0] - timedelta(seconds=1))
    tid = store.begin_turn(cid)
    return await run_turn(build_graph(), turn_id=tid, conversation_id=cid, context=ctx), provider


async def test_graph_uses_public_documents_and_cites_page():
    docs = FakeDocs([DOC])
    r, provider = await _turn(
        docs,
        "vào phòng thi được mang những gì?",
        lambda m: "Bạn được mang bút, thước và máy tính cầm tay [1].",
    )
    assert docs.calls == [frozenset({"public"})]
    assert "Quy chế thi tốt nghiệp THPT (trang 12)" in str(provider.calls[-1][0].content)
    sent = "\n".join(r["reply_parts"])
    assert "[1] Quy chế thi tốt nghiệp THPT (trang 12): Tài liệu: quy-che.pdf" in sent


async def test_graph_survives_document_search_failure():
    r, _ = await _turn(
        FakeDocs(fail=True), "vào phòng thi được mang gì?", lambda m: "Mình chưa có nguồn cho câu này."
    )
    assert r["action"] == "reply" and "doc_search_failed" in r["check_flags"]


async def test_internal_assistant_answers_with_sources_and_says_not_found():
    from app.rag.assistant import NOT_FOUND, ask

    s = get_settings()
    llm = FakeProvider(
        lambda m: "- Được mang bút viết, thước kẻ [1].\n- Máy tính cầm tay không có chức năng soạn thảo [1]."
    )
    res = await ask("thí sinh được mang gì vào phòng thi", settings=s, llm=llm, docs=FakeDocs([DOC]))
    assert (
        "Nguồn:\n[1] Quy chế thi tốt nghiệp THPT (trang 12)" in res["answer"] and res["sources"][0]["n"] == 1
    )
    empty_llm = FakeProvider(lambda m: "không nên gọi")
    res2 = await ask("lương giám đốc bao nhiêu", settings=s, llm=empty_llm, docs=FakeDocs([]))
    assert res2["answer"] == NOT_FOUND and empty_llm.calls == []


# ------------------------------------------------------------------------------------------ endpoint embeddings
def test_embeddings_endpoint_contract_and_limits():
    from fastapi.testclient import TestClient

    from serving.server import MAX_EMBED_INPUTS, create_app

    class FakeEmb:
        model_id = "fake-emb"

        def embed(self, texts):  # type: ignore[no-untyped-def]
            return [[0.1, 0.2] for _ in texts], 7

    class FakeRuntime:
        base_model_id, revision, adapter_id, adapter_method = "m", "r", None, None

        def vram_gb(self) -> float:
            return 0.0

    key = get_settings().model_server_api_key.get_secret_value()
    client = TestClient(create_app(FakeRuntime(), "m", FakeEmb()))  # type: ignore[arg-type]
    h = {"Authorization": f"Bearer {key}"}
    r = client.post("/v1/embeddings", json={"input": ["a", "b"]}, headers=h)
    assert r.status_code == 200 and [d["index"] for d in r.json()["data"]] == [0, 1]
    assert (
        client.post("/v1/embeddings", json={"input": ["a"] * (MAX_EMBED_INPUTS + 1)}, headers=h).status_code
        == 400
    )
    assert client.post("/v1/embeddings", json={"input": "a"}).status_code == 401


@pytest.mark.db
async def test_admin_ask_and_docs_require_auth(db):
    import httpx

    from app.main import create_app

    app = create_app()
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as c:
        assert (await c.post("/admin/ask", json={"question": "abc"})).status_code == 401
        assert (await c.get("/admin/docs")).status_code == 401
        key = get_settings().admin_api_key.get_secret_value()
        r = await c.get("/admin/docs", headers={"Authorization": f"Bearer {key}"})
        assert r.status_code == 200 and r.json() == []
