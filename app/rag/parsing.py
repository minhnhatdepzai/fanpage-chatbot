"""Đọc tài liệu (PDF, Word .docx, .txt/.md) -> trang văn bản sạch -> đoạn (chunk) có số trang để trích dẫn.

- PDF: pypdfium2 (PDFium, giấy phép Apache/BSD) trích lớp chữ; trang không có chữ (bản scan) trả về ảnh để OCR.
- Word: python-docx, gồm đoạn văn và bảng (mỗi hàng thành một dòng "ô | ô").
- Làm sạch: chuẩn hóa Unicode NFC (giữ dấu tiếng Việt), nối từ bị ngắt dòng bằng gạch nối, gộp khoảng trắng, bỏ dòng
  lặp lại ở đầu/cuối nhiều trang (header/footer) và dòng chỉ có số trang.
- Chia đoạn theo ranh giới câu/đoạn, có chồng lấp để không mất ngữ cảnh ở chỗ cắt.
Không chạy macro/script nào trong tài liệu.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SUPPORTED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
MIN_TEXT_CHARS_PER_PAGE = 25  # ít hơn -> coi là trang scan, cần OCR


class DocumentError(ValueError):
    pass


@dataclass
class Page:
    number: int  # bắt đầu từ 1
    text: str
    image: Any = None  # PIL.Image khi trang cần OCR


@dataclass
class ParsedDocument:
    mime: str
    pages: list[Page]
    sha256: str
    title_hint: str = ""
    meta: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class Chunk:
    index: int
    text: str
    page_from: int
    page_to: int


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean_text(text: str) -> str:
    t = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", ""))
    t = t.replace("­", "")  # soft hyphen
    # gạch nối ở cuối dòng: nối liền nhưng GIỮ gạch ("COVID-\n19" -> "COVID-19"); tiếng Việt không ghép âm tiết
    t = re.sub(r"(\w)-\n(\w)", r"\1-\2", t)
    t = re.sub(r"[ \t ]+", " ", t)
    t = re.sub(r" *\n *", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _strip_repeated_lines(pages: list[Page]) -> None:
    """Bỏ header/footer: dòng (đã chuẩn hóa số) xuất hiện ở đầu/cuối >= 60% số trang (khi có >= 3 trang)."""
    if len(pages) < 3:
        return

    def key(line: str) -> str:
        return re.sub(r"\d+", "#", line.strip().lower())

    edges: Counter[str] = Counter()
    for p in pages:
        lines = [x for x in p.text.split("\n") if x.strip()]
        for line in set(lines[:2] + lines[-2:]):
            edges[key(line)] += 1
    repeated = {k for k, c in edges.items() if c >= 0.6 * len(pages) and k}
    for p in pages:
        kept = [x for x in p.text.split("\n") if key(x) not in repeated]
        p.text = "\n".join(kept).strip()


def _drop_page_numbers(text: str) -> str:
    return "\n".join(
        x
        for x in text.split("\n")
        if not re.fullmatch(r"\s*(?:trang|page)?\s*[-–]?\s*\d{1,4}\s*[-–]?\s*(?:/\s*\d+)?\s*", x, re.I)
    )


def parse_pdf(data: bytes, max_pages: int, render_scale: float = 2.0) -> list[Page]:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise DocumentError(f"PDF không đọc được: {exc}") from exc
    try:
        if len(pdf) > max_pages:
            raise DocumentError(f"PDF có {len(pdf)} trang, vượt giới hạn {max_pages}")
        pages = []
        for i in range(len(pdf)):
            page = pdf[i]
            text = clean_text(page.get_textpage().get_text_range())
            image = None
            if len(text) < MIN_TEXT_CHARS_PER_PAGE:
                image = page.render(scale=render_scale).to_pil()
            pages.append(Page(i + 1, text, image))
        return pages
    finally:
        pdf.close()


def parse_docx(data: bytes) -> list[Page]:
    import io

    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx ném nhiều loại lỗi khác nhau cho tệp hỏng
        raise DocumentError(f"Word không đọc được: {type(exc).__name__}") from exc
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    # .docx không có khái niệm trang cố định -> coi cả tài liệu là "trang 1"
    return [Page(1, clean_text("\n".join(parts)))]


def parse_file(path: Path, *, max_bytes: int, max_pages: int) -> ParsedDocument:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(f"Định dạng {ext or '(không có)'} chưa hỗ trợ; dùng: {', '.join(SUPPORTED)}")
    size = path.stat().st_size
    if size > max_bytes:
        raise DocumentError(f"Tệp {size / 1e6:.1f} MB vượt giới hạn {max_bytes / 1e6:.0f} MB")
    data = path.read_bytes()
    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise DocumentError("Tệp đuôi .pdf nhưng nội dung không phải PDF")
        pages = parse_pdf(data, max_pages)
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise DocumentError("Tệp đuôi .docx nhưng nội dung không phải Word (zip)")
        pages = parse_docx(data)
    else:
        pages = [Page(1, clean_text(data.decode("utf-8", errors="replace")))]
    for p in pages:
        p.text = _drop_page_numbers(p.text)
    _strip_repeated_lines(pages)
    return ParsedDocument(SUPPORTED[ext], pages, sha256_bytes(data), title_hint=path.stem.replace("_", " "))


_SENT = re.compile(r"(?<=[.!?…;:])\s+|\n+")


def chunk_pages(pages: list[Page], *, max_chars: int = 900, overlap_chars: int = 150) -> list[Chunk]:
    """Gom câu thành đoạn <= max_chars (câu quá dài bị cắt cứng), chồng lấp ~overlap_chars; giữ khoảng trang."""
    units: list[tuple[str, int]] = []
    for p in pages:
        for sent in _SENT.split(p.text):
            sent = sent.strip()
            while len(sent) > max_chars:
                units.append((sent[:max_chars], p.number))
                sent = sent[max_chars - overlap_chars :]
            if sent:
                units.append((sent, p.number))
    chunks: list[Chunk] = []
    cur: list[tuple[str, int]] = []
    size = 0
    for unit in units:
        if cur and size + len(unit[0]) + 1 > max_chars:
            chunks.append(_mk(len(chunks), cur))
            tail: list[tuple[str, int]] = []
            tsize = 0
            for u in reversed(cur):  # chồng lấp: giữ lại vài câu cuối
                if tsize + len(u[0]) > overlap_chars:
                    break
                tail.insert(0, u)
                tsize += len(u[0]) + 1
            cur, size = tail, tsize
        cur.append(unit)
        size += len(unit[0]) + 1
    if cur:
        chunks.append(_mk(len(chunks), cur))
    seen: set[str] = set()
    out = []
    for c in chunks:  # bỏ đoạn trùng hệt nhau trong cùng tài liệu
        h = hashlib.sha1(c.text.encode(), usedforsecurity=False).hexdigest()
        if h not in seen:
            seen.add(h)
            out.append(Chunk(len(out), c.text, c.page_from, c.page_to))
    return out


def _mk(index: int, units: list[tuple[str, int]]) -> Chunk:
    return Chunk(index, " ".join(u[0] for u in units), units[0][1], units[-1][1])
